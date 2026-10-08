from __future__ import annotations

from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from auth import verify_credentials
from qrp_parser import QrpFormatError, ReportRow, parse_qrp


_COLUMNS = [
    "código CNES",
    "nome da unidade de saúde",
    "competência",
    "valor apresentado",
    "valor aprovado",
    "valor do serviço hospitalar",
    "valor do serviço profissional",
]


def _to_dataframe(rows: list[ReportRow]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "código CNES": row.cnes,
                "nome da unidade de saúde": row.unit_name,
                "competência": row.competence,
                "valor apresentado": row.presented,
                "valor aprovado": row.approved,
                "valor do serviço hospitalar": float(row.hospital_service),
                "valor do serviço profissional": float(row.professional_service),
            }
            for row in rows
        ],
        columns=_COLUMNS,
    )


def _make_excel(rows: list[ReportRow]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Aprovados por CNES"

    worksheet.append(_COLUMNS)
    for row in rows:
        worksheet.append(
            [
                row.cnes,
                row.unit_name,
                row.competence,
                row.presented,
                row.approved,
                float(row.hospital_service),
                float(row.professional_service),
            ]
        )

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["B"].width = 48
    worksheet.column_dimensions["C"].width = 14
    for column in ("D", "E"):
        worksheet.column_dimensions[column].width = 18
    for column in ("F", "G"):
        worksheet.column_dimensions[column].width = 30
    for row_number in range(2, worksheet.max_row + 1):
        for column in ("F", "G"):
            worksheet[f"{column}{row_number}"].number_format = '#,##0.00#'

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _show_login() -> None:
    try:
        auth_settings = st.secrets["auth"]
        configured_username = auth_settings["username"]
        password_salt = auth_settings["password_salt"]
        password_hash = auth_settings["password_hash"]
    except (KeyError, FileNotFoundError):
        st.error(
            "O login ainda não foi configurado. Adicione as credenciais "
            "hash em Streamlit secrets conforme o README."
        )
        st.stop()
    if not all(
        isinstance(value, str)
        for value in (configured_username, password_salt, password_hash)
    ):
        st.error(
            "A configuração de login deve conter usuário, salt e hash como "
            "textos. Revise Streamlit secrets."
        )
        st.stop()

    st.title("Acesso ao transformador SIHD2")
    with st.form("login_form"):
        username = st.text_input("Usuário")
        password = st.text_input("Senha", type="password")
        submitted = st.form_submit_button("Entrar")

    if submitted:
        if verify_credentials(
            username,
            password,
            configured_username,
            password_salt,
            password_hash,
        ):
            st.session_state["authenticated"] = True
            st.rerun()
        st.error("Usuário ou senha incorretos.")


st.set_page_config(page_title="Transformador de relatório SIHD2", page_icon="📊", layout="wide")

if not st.session_state.get("authenticated", False):
    _show_login()
    st.stop()

if st.sidebar.button("Sair"):
    st.session_state["authenticated"] = False
    st.rerun()

st.title("Transformador de relatório SIHD2")
st.write(
    "Envie o relatório mensal `.QRP` para extrair os dados por unidade de saúde "
    "e baixar uma planilha Excel."
)
st.caption(
    "Os valores apresentado e aprovado são as quantidades das colunas "
    "“Apres” e “Aprov” do relatório. Os serviços hospitalar e profissional "
    "são os respectivos totais, incluindo valores próprios e de terceiros."
)

uploaded_file = st.file_uploader("Selecione o arquivo .QRP", type=["qrp"])

if uploaded_file is not None:
    try:
        report_rows = parse_qrp(uploaded_file.getvalue())
    except QrpFormatError as error:
        st.error(str(error))
    else:
        dataframe = _to_dataframe(report_rows)
        competence = report_rows[0].competence
        st.success(
            f"Competência {competence}: {len(report_rows)} unidades extraídas."
        )
        st.dataframe(dataframe, use_container_width=True, hide_index=True)
        st.download_button(
            label="Baixar tabela em Excel",
            data=_make_excel(report_rows),
            file_name=f"aprovados_por_cnes_{competence.replace('/', '-')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
